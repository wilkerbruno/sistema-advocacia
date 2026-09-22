"""
Pedido explícito do usuário: "agora eu quero a opção de exportar que vai
baixar uma planilha com os dados" — baixa (.xlsx) exatamente o que está
sendo visto na tela do Financeiro (mesma aba "conta" e mesmos filtros de
status/natureza/unidade aplicados via querystring), tanto em "Caixa do
escritório" quanto em "Conta de terceiros" (pedido explícito do usuário:
"isso deve acontecer tanto em caixa do escritório quanto em conta de
terceiros").
"""
import io

import openpyxl
import pytest

from app.models import Lancamento
from tests.conftest import extrair_csrf


@pytest.fixture()
def cenario(app, empresa_basica, criar_usuario):
    unidade_id = empresa_basica["unidade_id"]
    admin_id = criar_usuario(unidade_id, "admin@exportar.com", papel="admin", nome="Admin Exportar")
    return dict(admin_id=admin_id, unidade_id=unidade_id)


def _criar_lancamento(client, unidade_id, **overrides):
    r = client.get("/financeiro/novo")
    token = extrair_csrf(r.data.decode("utf-8"))
    dados = dict(
        descricao="Lançamento de teste", valor="1000.00", natureza="receita",
        unidade_id=str(unidade_id), status="pendente", csrf_token=token,
    )
    dados.update(overrides)
    return client.post("/financeiro/novo", data=dados, follow_redirects=True)


def _planilha(response):
    return openpyxl.load_workbook(io.BytesIO(response.data))


def test_exportar_requer_login(client):
    r = client.get("/financeiro/exportar.xlsx")
    assert r.status_code in (302, 401)


def test_exportar_devolve_xlsx_valido_com_cabecalho(client, login, cenario):
    login("admin@exportar.com")
    _criar_lancamento(client, cenario["unidade_id"], descricao="Honorário de teste")

    r = client.get("/financeiro/exportar.xlsx")
    assert r.status_code == 200
    assert r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in r.headers.get("Content-Disposition", "")
    assert ".xlsx" in r.headers.get("Content-Disposition", "")

    wb = _planilha(r)
    ws = wb.active
    cabecalhos = [c.value for c in ws[1]]
    assert "Descrição" in cabecalhos
    assert "Data de pagamento" in cabecalhos
    assert "Status" in cabecalhos
    # Pedido explícito: coluna de data de pagamento ANTES da de status.
    assert cabecalhos.index("Data de pagamento") < cabecalhos.index("Status")

    linhas = list(ws.iter_rows(min_row=2, values_only=True))
    descricoes = [linha[cabecalhos.index("Descrição")] for linha in linhas]
    assert "Honorário de teste" in descricoes


def test_exportar_respeita_filtro_de_conta_terceiros(client, login, cenario):
    login("admin@exportar.com")
    _criar_lancamento(client, cenario["unidade_id"], descricao="Receita operacional normal")
    _criar_lancamento(client, cenario["unidade_id"], descricao="Depósito judicial de terceiro",
                       conta_terceiros="1")

    wb_operacional = _planilha(client.get("/financeiro/exportar.xlsx?conta=operacional"))
    ws = wb_operacional.active
    cabecalhos = [c.value for c in ws[1]]
    descricoes_operacional = [linha[cabecalhos.index("Descrição")]
                               for linha in ws.iter_rows(min_row=2, values_only=True)]
    assert "Receita operacional normal" in descricoes_operacional
    assert "Depósito judicial de terceiro" not in descricoes_operacional

    wb_terceiros = _planilha(client.get("/financeiro/exportar.xlsx?conta=terceiros"))
    ws2 = wb_terceiros.active
    descricoes_terceiros = [linha[cabecalhos.index("Descrição")]
                             for linha in ws2.iter_rows(min_row=2, values_only=True)]
    assert "Depósito judicial de terceiro" in descricoes_terceiros
    assert "Receita operacional normal" not in descricoes_terceiros


def test_exportar_respeita_filtro_de_status(client, login, cenario):
    login("admin@exportar.com")
    _criar_lancamento(client, cenario["unidade_id"], descricao="Pendente de teste", status="pendente")
    _criar_lancamento(client, cenario["unidade_id"], descricao="Pago de teste", status="pago")

    wb = _planilha(client.get("/financeiro/exportar.xlsx?status=pago"))
    ws = wb.active
    cabecalhos = [c.value for c in ws[1]]
    descricoes = [linha[cabecalhos.index("Descrição")] for linha in ws.iter_rows(min_row=2, values_only=True)]
    assert "Pago de teste" in descricoes
    assert "Pendente de teste" not in descricoes


def test_exportar_inclui_data_de_pagamento_e_valor_atualizado(client, login, cenario):
    login("admin@exportar.com")
    _criar_lancamento(
        client, cenario["unidade_id"], descricao="Já pago com data",
        status="pago", data_pagamento="2026-01-15",
    )
    _criar_lancamento(
        client, cenario["unidade_id"], descricao="Vencido com multa e juros",
        status="pendente", data_vencimento="2020-01-01",
        multa_tipo="valor", multa_valor="50.00", juros_tipo="dia", juros_valor="0.1",
    )

    wb = _planilha(client.get("/financeiro/exportar.xlsx"))
    ws = wb.active
    cabecalhos = [c.value for c in ws[1]]
    linhas = {linha[cabecalhos.index("Descrição")]: linha for linha in ws.iter_rows(min_row=2, values_only=True)}

    idx_pgto = cabecalhos.index("Data de pagamento")
    assert linhas["Já pago com data"][idx_pgto] == "15/01/2026"

    idx_valor_atualizado = cabecalhos.index("Valor atualizado (com multa/juros)")
    idx_multa = cabecalhos.index("Multa")
    idx_juros = cabecalhos.index("Juros de atraso")
    linha_vencida = linhas["Vencido com multa e juros"]
    assert linha_vencida[idx_valor_atualizado] is not None
    assert linha_vencida[idx_valor_atualizado] > 1000.0
    assert linha_vencida[idx_multa] == "R$ 50,00"
    assert "ao dia" in linha_vencida[idx_juros]
    assert linha_vencida[idx_juros].startswith("0.1")

    # Lançamento já pago não está "vencido" — não recebe valor atualizado
    # (openpyxl grava string vazia como célula vazia, lida de volta como None).
    assert linhas["Já pago com data"][idx_valor_atualizado] in (None, "")


def test_exportar_nome_do_arquivo_inclui_a_conta(client, login, cenario):
    login("admin@exportar.com")
    r = client.get("/financeiro/exportar.xlsx?conta=terceiros")
    assert "financeiro_terceiros_" in r.headers.get("Content-Disposition", "")
